# Rellena la plantilla "Application Material (for Individual Developers).xlsx" de Huawei Health
# Service Kit (la que corresponde a cuenta individual). Texto en negro, como pide la plantilla.
#   python docs/huawei/fill_individual_xlsx.py [captura_permisos.png]
import sys
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Font
from PIL import Image

HERE = Path(__file__).parent
OUT = HERE / "HomeTest_Huawei_Application_Material_Individual.xlsx"
wb = openpyxl.load_workbook(HERE / "template_individual.xlsx")
WRAP = Alignment(wrap_text=True, vertical="top")


def put(ws, cell, value):
    c = ws[cell]
    c.value = value
    c.font = Font(name=c.font.name or "Arial", size=c.font.size or 10, color="FF000000", bold=c.font.bold)
    c.alignment = WRAP


def image(ws, cell, path, width):
    im = Image.open(path)
    tmp = HERE / f"_xl_{Path(path).stem}.png"
    im.convert("RGB").save(tmp)
    xl = XlImage(str(tmp))
    xl.width, xl.height = width, int(width * im.height / im.width)
    ws.add_image(xl, cell)


# --- Hoja 1: Data Usage ---
ws = wb.worksheets[0]
ws._images = []  # quitar las imágenes de ejemplo
DONE = "2026-11-30"
rows = [
    ("Read step count data",
     "Read daily step totals (classified by day) and show them as a daily ring and 14-day / 1-month trends.",
     "After the user connects Huawei Health and authorises access, steps are shown in the user's own dashboard to track a personalised activity goal (e.g. 8,000 steps/day) next to their blood test results."),
    ("Read calories burned (active energy) data",
     "Read daily active calories (classified by day) and show them as a daily ring and 14-day / 1-month trends.",
     "Give users a simple view of their daily energy expenditure alongside their other health indicators."),
    ("Read distance data",
     "Read daily distance (classified by day) to complete the activity summary.",
     "Contextualise daily activity (steps vs. distance) for the user."),
    ("Historical data: one month (read-only)",
     "On first connection, read up to the previous month of steps, calories and distance.",
     "So the user sees a baseline and recent trends from day one instead of empty charts. Not used for any other purpose."),
]
for i, (scope, req, scen) in enumerate(rows):
    r = 3 + i
    put(ws, f"A{r}", i + 1)
    put(ws, f"B{r}", scope)
    put(ws, f"C{r}", req)
    put(ws, f"D{r}", scen)
    put(ws, f"E{r}", "Not supported (Android only for now)")
    put(ws, f"G{r}", DONE)
    ws.row_dimensions[r].height = 120
put(ws, "A7", None); ws["A7"].value = None
put(ws, "B7", None)
put(ws, "F3", "Android mobile app")
put(ws, "H3", "Today screen (see image). Sample data.")
put(ws, "I3", "App ID: 119155453\nPackage: com.astoramor.apptestsmedicos\nScopes selected in the console: Step, Calories, Distance (read) + one month of history.")
image(ws, "H4", HERE / "today_2.jpg", 170)
if len(sys.argv) > 1:
    image(ws, "I5", sys.argv[1], 250)
put(ws, "J3", ws["J3"].value)  # remarks de la plantilla, en negro

# --- Hoja 2: App Info ---
ws = wb.worksheets[1]
put(ws, "B2", "HomeTest is a preventive-health membership app (Spain, EU). Users record and understand their blood "
    "test results, home measurements, wellbeing check-ins and daily activity from their wearables. It shows trends "
    "and informational wellbeing recommendations (it does not diagnose or treat) and lets users share selected data "
    "with the professionals they choose. Huawei Health data (steps, calories, distance) is shown only in the user's "
    "own dashboard. Data stored in the EU, never sold or used for advertising. Applicant: individual developer "
    "Astor García Amor.")
put(ws, "B3", "Today screen of the HomeTest app (sample data): daily activity rings (steps, active calories).")
image(ws, "C3", HERE / "today_screenshot.png", 520)
put(ws, "B4", "Health and fitness / preventive healthcare (consumer)")
put(ws, "B5", "Show users their daily activity next to their at-home blood test results, as part of a personalised "
    "preventive-health plan. Demo / pre-release; commercial release planned after the company is incorporated in Spain.")
put(ws, "B6", "Step count, active calories and distance (read-only, one month of history) are shown only to the user "
    "in their own dashboard: a daily ring on the Today screen and 14-day / 1-month trend charts, to track a "
    "personalised activity goal. No write scopes, no sharing with third parties without explicit consent.")
ws.column_dimensions["B"].width = 90
for r in range(2, 7):
    ws.row_dimensions[r].height = 90

# --- Hoja 3: Self-Check ---
ws = wb.worksheets[2]
put(ws, "C2", "Yes. The app is to be released in Spain (European Union).")
put(ws, "C3", "Demo (pre-release testing; commercial use after company incorporation).")
put(ws, "C4", "No.")
put(ws, "C5", "Android mobile apps.")
put(ws, "C6", "RESTful API (cloud-side, OAuth 2.0 with Huawei ID; tokens stored server-side).")
put(ws, "C7", "Yes.")
put(ws, "C8", "No. Wellbeing information only; no diagnosis or treatment.")
put(ws, "C9", "Pilot: fewer than 1,000 users in the first year; TPS < 1.")
put(ws, "A10", ws["A10"].value)

wb.save(OUT)
for f in HERE.glob("_xl_*.png"):
    pass  # openpyxl lee las imágenes al guardar; se pueden borrar después
print(OUT)
